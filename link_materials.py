"""Link catalog records to actual materials in genshin_v2.db.

Copy beside Genshin_Tracker.xlsm and genshin_v2.db. Run with project Python.
Source lookup gaps are recorded in material_mapping_issues.
"""
import sqlite3
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

folder = Path(__file__).resolve().parent
db_path = folder / 'genshin_v2.db'
workbook_path = folder / 'Genshin_Tracker.xlsm'
if not db_path.is_file():
    raise SystemExit('genshin_v2.db not found beside this script')

book = load_workbook(workbook_path, read_only=True, data_only=True)
enemy = book['Enemy Drops Table']
talents = book['Character Talent Mats']
ascension = book['Weapon Ascension Mats']

common = {r[0]: r[1:4] for r in enemy.iter_rows(min_row=3, max_row=18, min_col=1, max_col=4, values_only=True) if r[0]}
elite = {r[0]: r[1:4] for r in enemy.iter_rows(min_row=22, max_row=47, min_col=1, max_col=4, values_only=True) if r[0]}
weapon_ascension = {r[0]: r[1:5] for r in ascension.iter_rows(min_row=3, max_row=24, min_col=1, max_col=5, values_only=True) if r[0]}
talent_books = {r[0]: r[1] for r in talents.iter_rows(min_row=3, max_row=24, min_col=1, max_col=2, values_only=True) if r[0]}
weekly = {r[0]: r[1] for r in talents.iter_rows(min_row=3, max_row=39, min_col=4, max_col=5, values_only=True) if r[0]}

db = sqlite3.connect(db_path)
db.execute('PRAGMA foreign_keys = ON')
try:
    db.executescript('''
        CREATE TABLE IF NOT EXISTS character_materials (
            character_id INTEGER NOT NULL REFERENCES characters(id),
            material_role TEXT NOT NULL,
            tier INTEGER NOT NULL,
            material_id INTEGER NOT NULL REFERENCES materials(id),
            PRIMARY KEY (character_id, material_role, tier)
        );
        CREATE TABLE IF NOT EXISTS weapon_materials (
            weapon_id INTEGER NOT NULL REFERENCES weapons(id),
            material_role TEXT NOT NULL,
            tier INTEGER NOT NULL,
            material_id INTEGER NOT NULL REFERENCES materials(id),
            PRIMARY KEY (weapon_id, material_role, tier)
        );
        CREATE TABLE IF NOT EXISTS material_mapping_issues (
            entity_kind TEXT NOT NULL,
            entity_id INTEGER NOT NULL,
            material_role TEXT NOT NULL,
            missing_family TEXT NOT NULL,
            PRIMARY KEY (entity_kind, entity_id, material_role)
        );
    ''')
    db.execute('DELETE FROM character_materials')
    db.execute('DELETE FROM weapon_materials')
    db.execute('DELETE FROM material_mapping_issues')

    added = set()
    def link(table, entity_id, role, tier, name, category):
        if not name:
            raise ValueError(f'Empty material name for {table} {entity_id}, {role}, tier {tier}')
        row = db.execute('SELECT id FROM materials WHERE name = ?', (name,)).fetchone()
        if row is None:
            db.execute('INSERT INTO materials (name, category) VALUES (?, ?)', (name, category))
            row = db.execute('SELECT id FROM materials WHERE name = ?', (name,)).fetchone()
            db.execute('INSERT INTO inventory (material_id, quantity) VALUES (?, 0)', (row[0],))
            added.add(name)
        id_col = 'character_id' if table == 'character_materials' else 'weapon_id'
        db.execute(f'INSERT INTO {table} ({id_col}, material_role, tier, material_id) VALUES (?, ?, ?, ?)',
                   (entity_id, role, tier, row[0]))

    def link_family(table, entity_id, role, family, lookup, first_tier, category):
        names = lookup.get(family)
        if names is None or any(not n for n in names):
            kind = 'character' if table == 'character_materials' else 'weapon'
            db.execute('INSERT INTO material_mapping_issues VALUES (?, ?, ?, ?)', (kind, entity_id, role, str(family)))
            return
        for offset, name in enumerate(names):
            link(table, entity_id, role, first_tier + offset, name, category)

    characters = db.execute('''SELECT id, name, gem_family, boss_material,
        common_drop_family, local_specialty, talent_book_family, weekly_boss_material
        FROM characters''').fetchall()
    for cid, name, gem, boss, common_family, local, talent_family, weekly_key in characters:
        table = 'character_materials'
        for tier, suffix in ((2,'Sliver'),(3,'Fragment'),(4,'Chunk'),(5,'Gemstone')):
            link(table,cid,'gem',tier,f'{gem} {suffix}','ascension_gem')
        link(table,cid,'boss_drop',0,boss,'character_boss_drop')
        link_family(table,cid,'common_drop',common_family,common,1,'common_drop')
        link(table,cid,'local_specialty',0,local,'local_specialty')
        book_name = talent_books.get(talent_family)
        if book_name is None:
            db.execute('INSERT INTO material_mapping_issues VALUES (?, ?, ?, ?)', ('character',cid,'talent_book',str(talent_family)))
        else:
            for tier,prefix in ((2,'Teachings of '),(3,'Guide to '),(4,'Philosophies of ')):
                link(table,cid,'talent_book',tier,prefix+book_name,'talent_book')
        weekly_name=weekly.get(weekly_key)
        if weekly_name is None:
            db.execute('INSERT INTO material_mapping_issues VALUES (?, ?, ?, ?)', ('character',cid,'weekly_boss_drop',str(weekly_key)))
        else:
            link(table,cid,'weekly_boss_drop',0,weekly_name,'weekly_boss_drop')
        link(table,cid,'crown',0,'Crown of Insight','talent_special')
        link(table,cid,'mora',0,'Mora','currency')

    weapons = db.execute('''SELECT id, name, ascension_family,
        common_drop_family, elite_drop_family FROM weapons''').fetchall()
    for wid, name, asc_family, common_family, elite_family in weapons:
        table = 'weapon_materials'
        link_family(table,wid,'weapon_ascension',asc_family,weapon_ascension,2,'weapon_ascension')
        link_family(table,wid,'common_drop',common_family,common,1,'common_drop')
        link_family(table,wid,'elite_drop',elite_family,elite,2,'elite_drop')
        link(table,wid,'mora',0,'Mora','currency')
    db.commit()

    for table in ('character_materials','weapon_materials'):
        print(f'{table}:',db.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0])
    issues=db.execute('''SELECT i.entity_kind, i.material_role, i.missing_family, COUNT(*)
        FROM material_mapping_issues i GROUP BY i.entity_kind,i.material_role,i.missing_family
        ORDER BY i.entity_kind,i.material_role,i.missing_family''').fetchall()
    print('new materials with zero owned:', ', '.join(sorted(added)) or 'none')
    print('incomplete characters:',db.execute("SELECT COUNT(DISTINCT entity_id) FROM material_mapping_issues WHERE entity_kind='character'").fetchone()[0])
    print('incomplete weapons:',db.execute("SELECT COUNT(DISTINCT entity_id) FROM material_mapping_issues WHERE entity_kind='weapon'").fetchone()[0])
    print('missing family links (entity, role, family, records):')
    for issue in issues:print('  ',*issue)
finally:
    db.close()
    book.close()
