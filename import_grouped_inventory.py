import sqlite3
from pathlib import Path

from openpyxl import load_workbook

folder = Path(__file__).parent
workbook = load_workbook(
    folder / "Genshin_Tracker.xlsm",
    read_only=True,
    data_only=True
)
overview = workbook["Overview"]

talent_sheet = workbook["Character Talent Mats"]
weapon_sheet = workbook["Weapon Ascension Mats"]
enemy_sheet = workbook["Enemy Drops Table"]

talent_names = {
    row[0]: row[1]
    for row in talent_sheet.iter_rows(
        min_row=3, max_row=24, min_col=1, max_col=2, values_only=True
    )
    if row[0]
}

weapon_names = {
    row[0]: row[1:5]
    for row in weapon_sheet.iter_rows(
        min_row=3, max_row=24, min_col=1, max_col=5, values_only=True
    )
    if row[0]
}

common_names = {
    row[0]: row[1:4]
    for row in enemy_sheet.iter_rows(
        min_row=3, max_row=18, min_col=1, max_col=4, values_only=True
    )
    if row[0]
}

elite_names = {
    row[0]: row[1:4]
    for row in enemy_sheet.iter_rows(
        min_row=22, max_row=47, min_col=1, max_col=4, values_only=True
    )
    if row[0]
}

database = sqlite3.connect(folder / "genshin_v2.db")
database.execute("PRAGMA foreign_keys = ON")

try:
    def save_item(name, category, quantity):
        database.execute("""
            INSERT INTO materials (name, category)
            VALUES (?, ?)
            ON CONFLICT(name) DO UPDATE SET category = excluded.category
        """, (name, category))

        material_id = database.execute(
            "SELECT id FROM materials WHERE name = ?",
            (name,)
        ).fetchone()[0]

        database.execute("""
            INSERT INTO inventory (material_id, quantity)
            VALUES (?, ?)
            ON CONFLICT(material_id)
            DO UPDATE SET quantity = excluded.quantity
        """, (material_id, int(quantity)))

    sections = [
        # Category, family column, tier column, owned column, first tier, lookup
        ("talent_book", "E", "F", "H", 2, talent_names),
        ("weapon_ascension", "L", "M", "O", 2, weapon_names),
        ("elite_drop", "S", "T", "V", 2, elite_names),
        ("common_drop", "Z", "AA", "AC", 1, common_names),
    ]

    imported = 0

    for category, family_col, tier_col, owned_col, first_tier, lookup in sections:
        family = None

        for row_number in range(3, 88):
            heading = overview[f"{family_col}{row_number}"].value
            if isinstance(heading, str):
                family = heading

            tier = overview[f"{tier_col}{row_number}"].value
            owned = overview[f"{owned_col}{row_number}"].value

            if not isinstance(tier, int) or not isinstance(owned, (int, float)):
                continue

            if family not in lookup:
                raise ValueError(
                    f"Unknown {category} family {family!r} on row {row_number}"
                )

            if category == "talent_book":
                prefixes = ("Teachings of ", "Guide to ", "Philosophies of ")
                name = prefixes[tier - 2] + lookup[family]
            else:
                name = lookup[family][tier - first_tier]

            if not name:
                raise ValueError(f"Missing material name on row {row_number}")

            save_item(name, category, owned)
            imported += 1

    database.commit()

    total = database.execute(
        "SELECT COUNT(*) FROM inventory"
    ).fetchone()[0]

    print(f"grouped material rows imported: {imported}")
    print(f"total inventory entries: {total}")

finally:
    database.close()
    workbook.close()