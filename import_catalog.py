import sqlite3
from pathlib import Path

from openpyxl import load_workbook

folder = Path(__file__).parent
workbook_path = folder / "Genshin_Tracker.xlsm"
database_path = folder / "genshin_v2.db"

workbook = load_workbook(workbook_path, read_only=True, data_only=True)
database = sqlite3.connect(database_path)

try:
    database.execute("""
        CREATE TABLE IF NOT EXISTS characters (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL UNIQUE,
            element TEXT,
            gem_family TEXT,
            boss_material TEXT,
            common_drop_family TEXT,
            local_specialty TEXT,
            talent_book_family TEXT,
            weekly_boss_material TEXT
        )
    """)

    database.execute("""
        CREATE TABLE IF NOT EXISTS weapons (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL UNIQUE,
            weapon_type TEXT,
            rarity INTEGER,
            ascension_family TEXT,
            common_drop_family TEXT,
            elite_drop_family TEXT
        )
    """)

    for row in workbook["Character Data Table"].iter_rows(
        min_row=3, min_col=2, max_col=9, values_only=True
    ):
        if not row[0]:
            continue

        database.execute("""
            INSERT OR IGNORE INTO characters (
                name, element, gem_family, boss_material,
                common_drop_family, local_specialty,
                talent_book_family, weekly_boss_material
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, row)

    for row in workbook["Weapon Data Table"].iter_rows(
        min_row=3, min_col=2, max_col=7, values_only=True
    ):
        if not row[0]:
            continue

        database.execute("""
            INSERT OR IGNORE INTO weapons (
                name, weapon_type, rarity, ascension_family,
                common_drop_family, elite_drop_family
            ) VALUES (?, ?, ?, ?, ?, ?)
        """, row)

    database.commit()

    for table in ("characters", "weapons"):
        count = database.execute(
            f"SELECT COUNT(*) FROM {table}"
        ).fetchone()[0]
        print(f"{table}: {count}")

finally:
    database.close()
    workbook.close()