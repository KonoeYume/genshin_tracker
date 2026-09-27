import sqlite3
from pathlib import Path

from openpyxl import load_workbook

folder = Path(__file__).parent
workbook = load_workbook(
    folder / "Genshin_Tracker.xlsm",
    read_only=True,
    data_only=True
)
database = sqlite3.connect(folder / "genshin_v2.db")
database.execute("PRAGMA foreign_keys = ON")

try:
    sheet = workbook["Character Data Table"]

    local_specialties = {
        row[0]
        for row in sheet.iter_rows(
            min_row=3, min_col=7, max_col=7, values_only=True
        )
        if isinstance(row[0], str) and row[0].strip()
    }

    items = [
        (name, "local_specialty")
        for name in sorted(local_specialties)
    ]
    items.append(("Crown of Insight", "talent_special"))

    for name, category in items:
        database.execute("""
            INSERT OR IGNORE INTO materials (name, category)
            VALUES (?, ?)
        """, (name, category))

        material_id = database.execute(
            "SELECT id FROM materials WHERE name = ?",
            (name,)
        ).fetchone()[0]

        database.execute("""
            INSERT OR IGNORE INTO inventory (material_id, quantity)
            VALUES (?, 0)
        """, (material_id,))

    database.commit()

    total = database.execute(
        "SELECT COUNT(*) FROM inventory"
    ).fetchone()[0]

    print(f"local specialties found: {len(local_specialties)}")
    print(f"total inventory entries: {total}")

finally:
    database.close()
    workbook.close()