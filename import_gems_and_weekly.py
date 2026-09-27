import sqlite3
from pathlib import Path

from openpyxl import load_workbook

folder = Path(__file__).parent
workbook = load_workbook(
    folder / "Genshin_Tracker.xlsm",
    read_only=True,
    data_only=True
)
overview = workbook["Overview"]
talent_sheet = workbook["Character Talent Mats"]

weekly_names = {
    row[0]: row[1]
    for row in talent_sheet.iter_rows(
        min_row=3, max_row=39, min_col=4, max_col=5, values_only=True
    )
    if row[0]
}

database = sqlite3.connect(folder / "genshin_v2.db")
database.execute("PRAGMA foreign_keys = ON")

try:
    def save_item(name, category, quantity):
        if not isinstance(quantity, (int, float)):
            raise ValueError(f"Invalid quantity for {name}: {quantity}")

        database.execute("""
            INSERT INTO materials (name, category)
            VALUES (?, ?)
            ON CONFLICT(name) DO UPDATE SET category = excluded.category
        """, (name, category))

        material_id = database.execute(
            "SELECT id FROM materials WHERE name = ?",
            (name,)
        ).fetchone()[0]

        database.execute("""
            INSERT INTO inventory (material_id, quantity)
            VALUES (?, ?)
            ON CONFLICT(material_id)
            DO UPDATE SET quantity = excluded.quantity
        """, (material_id, int(quantity)))

    suffixes = {
        2: "Sliver",
        3: "Fragment",
        4: "Chunk",
        5: "Gemstone",
    }

    gem_family = None
    gem_count = 0

    for row_number in range(3, 31):
        heading = overview[f"AG{row_number}"].value
        if isinstance(heading, str):
            gem_family = heading

        tier = overview[f"AH{row_number}"].value
        owned = overview[f"AJ{row_number}"].value

        if not isinstance(tier, int):
            continue

        name = f"{gem_family} {suffixes[tier]}"
        save_item(name, "ascension_gem", owned)
        gem_count += 1

    weekly_count = 0

    for row_number in range(4, 16):
        boss = overview[f"AS{row_number}"].value

        for tier, column in enumerate(("AY", "AZ", "BA"), start=1):
            key = f"{boss} {tier}"
            name = weekly_names.get(key)

            if name is None:
                raise ValueError(f"No weekly material for {key}")

            owned = overview[f"{column}{row_number}"].value
            save_item(name, "weekly_boss_drop", owned)
            weekly_count += 1

    database.commit()

    total = database.execute(
        "SELECT COUNT(*) FROM inventory"
    ).fetchone()[0]

    print(f"gems imported: {gem_count}")
    print(f"weekly boss drops imported: {weekly_count}")
    print(f"total inventory entries: {total}")

finally:
    database.close()
    workbook.close()