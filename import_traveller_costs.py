"""Import Traveller material rules from Genshin_Tracker.xlsm into genshin_v2.db.

Run once beside both files. Safe to rerun; progress and inventory are preserved.
Cryo has three progress placeholders but no workbook rules yet.
"""
import re
import sqlite3
from pathlib import Path
from openpyxl import load_workbook

FOLDER = Path(__file__).resolve().parent
ELEMENTS = ('Anemo', 'Geo', 'Electro', 'Dendro', 'Hydro', 'Pyro')
# Workbook formulas contain IF(AND(B$4<5,B$5>=5),6,0), for example.
MILESTONE = re.compile(r'\$4<(\d+),[BEH]\$5>=\1\),(\d+),0\)')


def main():
    book_path, db_path = FOLDER / 'Genshin_Tracker.xlsm', FOLDER / 'genshin_v2.db'
    if not book_path.is_file() or not db_path.is_file():
        raise SystemExit('Place this script beside Genshin_Tracker.xlsm and genshin_v2.db')
    formulas = load_workbook(book_path, data_only=False, read_only=False)
    cached = load_workbook(book_path, data_only=True, read_only=True)
    db = sqlite3.connect(db_path)
    try:
        db.execute('PRAGMA foreign_keys = ON')
        names = {name: mid for mid, name in db.execute('SELECT id, name FROM materials')}
        standard = {(level, role, tier): quantity for level, role, tier, quantity in
                    db.execute('SELECT talent_level, material_role, tier, cumulative_quantity FROM talent_costs')}
        records = []
        book_totals = {}
        for element in ELEMENTS:
            sf, sc = formulas['Traveller '+element], cached['Traveller '+element]
            for slot, (amount_col, name_col) in enumerate((('B','A'),('E','D'),('H','G')), 1):
                book_totals[element, slot] = 0
                for row in range(6,21):
                    name = sc[f'{name_col}{row}'].value
                    if not name:
                        raise ValueError(f'Missing cached material name: {element} slot {slot} row {row}')
                    if name not in names:
                        # These workbook materials were absent from the earlier general catalog.
                        if name not in ('Cornerstone of Stars and Flames',):
                            raise ValueError(f'Material not imported: {name}')
                        db.execute('INSERT INTO materials(name,category) VALUES (?,?)',
                                   (name,'weekly_boss_drop'))
                        names[name] = db.execute('SELECT last_insert_rowid()').fetchone()[0]
                        db.execute('INSERT INTO inventory(material_id,quantity) VALUES (?,0)',(names[name],))
                    formula = sf[f'{amount_col}{row}'].value
                    if row <= 14:
                        if not isinstance(formula,str) or not formula.startswith('='):
                            raise ValueError(f'Unexpected book formula at {element} {amount_col}{row}')
                        pairs = [(int(level),int(amount)) for level,amount in MILESTONE.findall(formula)]
                        if formula != '=0' and not pairs:
                            raise ValueError(f'Cannot read book milestones: {element} {amount_col}{row}: {formula}')
                        by_level = dict(pairs)
                        if len(by_level) != len(pairs):
                            raise ValueError(f'Duplicate milestone: {element} {amount_col}{row}')
                        for level, amount in by_level.items():
                            records.append((element,slot,level,names[name],amount))
                            book_totals[element, slot] += amount
                    else:
                        role,tier = ({15:('weekly_boss_drop',0),16:('crown',0),
                                      17:('common_drop',1),18:('common_drop',2),
                                      19:('common_drop',3),20:('mora',0)})[row]
                        for level in range(2,11):
                            quantity = standard[(level,role,tier)] - standard[(level-1,role,tier)]
                            if quantity:
                                records.append((element,slot,level,names[name],quantity))
        # Validate that each Traveller slot's total book count matches standard talents.
        for element in ELEMENTS:
            for slot in (1,2,3):
                books = book_totals[element,slot]
                expected = sum(standard[(10,'talent_book',tier)] for tier in (2,3,4))
                if books != expected:
                    raise ValueError(f'Book total mismatch for {element} slot {slot}: {books} vs {expected}')
        db.execute('''CREATE TABLE IF NOT EXISTS traveller_talent_costs (
            element TEXT NOT NULL, talent_slot INTEGER NOT NULL,
            talent_level INTEGER NOT NULL, material_id INTEGER NOT NULL REFERENCES materials(id),
            quantity INTEGER NOT NULL CHECK(quantity > 0),
            PRIMARY KEY(element,talent_slot,talent_level,material_id))''')
        db.execute('DELETE FROM traveller_talent_costs')
        db.executemany('INSERT INTO traveller_talent_costs VALUES(?,?,?,?,?)',records)
        db.execute('''CREATE TABLE IF NOT EXISTS traveller_level_materials (
            material_role TEXT NOT NULL, tier INTEGER NOT NULL,
            material_id INTEGER NOT NULL REFERENCES materials(id),
            PRIMARY KEY(material_role,tier))''')
        db.execute('DELETE FROM traveller_level_materials')
        level_sheet = cached['Traveller Level']
        for row, role, tier, category in (
            (12,'gem',2,'ascension_gem'),(13,'gem',3,'ascension_gem'),
            (14,'gem',4,'ascension_gem'),(15,'gem',5,'ascension_gem'),
            (16,'common_drop',1,'common_drop'),(17,'common_drop',2,'common_drop'),
            (18,'common_drop',3,'common_drop'),(19,'local_specialty',0,'local_specialty'),
            (20,'mora',0,'currency')):
            name = level_sheet[f'B{row}'].value
            if not name:
                raise ValueError(f'Missing Traveller level material at B{row}')
            if name not in names:
                if role != 'gem':
                    raise ValueError(f'Material not imported: {name}')
                db.execute('INSERT INTO materials(name,category) VALUES (?,?)',(name,category))
                names[name] = db.execute('SELECT last_insert_rowid()').fetchone()[0]
                db.execute('INSERT INTO inventory(material_id,quantity) VALUES (?,0)',(names[name],))
            db.execute('INSERT INTO traveller_level_materials VALUES (?,?,?)',
                       (role,tier,names[name]))
        db.commit()
        print(f'Traveller talent cost entries: {len(records)} across 6 elements and 18 talents')
        print('Cryo remains a placeholder with no cost rules.')
    except Exception:
        db.rollback()
        raise
    finally:
        db.close(); cached.close(); formulas.close()

if __name__ == '__main__':
    main()
